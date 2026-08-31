# -*- coding: utf-8 -*-
"""从工作项 JSON 生成工作量 Excel（导出结果格式，14 列）。"""
import json

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def build(cfg, items, out_path, status=None, updated_at=None, created_at=None):
    """items: [{"title","hours","start","end","description"}]
    生成 14 列「导出结果」格式工作簿。
    """
    status = status or cfg.get("initial_status", "新建")
    updated_at = updated_at or ""
    created_at = created_at or ""
    columns = cfg.get("excel_columns", [
        "编号", "标题", "工作项类型", "状态", "指派给", "更新时间",
        "计划开始时间", "实际完成时间", "初始估计", "创建人", "创建时间",
        "团队", "详细说明", "任务类型"])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "导出结果"
    hfont = Font(name="宋体", size=14, bold=True)
    hfill = PatternFill(patternType="solid", fgColor="D9D9D9")
    thin = Side(style="thin")
    hborder = Border(left=thin, right=thin, top=thin, bottom=thin)
    halign = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for j, h in enumerate(columns, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = hfont; c.fill = hfill; c.border = hborder; c.alignment = halign

    for i, it in enumerate(items, start=2):
        vals = {
            "编号": it.get("id", ""),
            "标题": it["title"],
            "工作项类型": cfg.get("work_item_type", "任务"),
            "状态": status,
            "指派给": f"{cfg.get('assignee_name','')}{cfg.get('assignee_emp_no','')}",
            "更新时间": updated_at,
            "计划开始时间": f"{it.get('start','')} 09:00:00" if it.get("start") else "",
            "实际完成时间": f"{it.get('end','')} 18:00:00" if it.get("end") else "",
            "初始估计": str(it.get("hours", "")),
            "创建人": f"{cfg.get('assignee_name','')}{cfg.get('assignee_emp_no','')}",
            "创建时间": created_at,
            "团队": cfg.get("team_name", ""),
            "详细说明": it.get("description", ""),
            "任务类型": cfg.get("task_type", "开发"),
        }
        for j, h in enumerate(columns, start=1):
            c = ws.cell(row=i, column=j, value=vals.get(h, ""))
            c.alignment = Alignment(vertical="center", wrap_text=True)
    widths = {"A": 24, "B": 46, "C": 12, "D": 10, "E": 22, "F": 20, "G": 22,
              "H": 22, "I": 12, "J": 22, "K": 20, "L": 32, "M": 90, "N": 12}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(out_path)
    return {"items": len(items), "status": status, "out": out_path}


def load_items(path):
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    return data.get("work_items", data if isinstance(data, list) else [])
